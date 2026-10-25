#!/usr/bin/env python3
"""Month-end billing for a Jira Server/Data Center + Tempo project, run from inside the project's Product Brain hub.
The project's config is the hub's brain.config.json > "finance" (or --config FILE); work files go to
<hub>/.work/finance/<id>/<month>/.

  check   [--name N --jira-url U --keys K --shared S] [--finance XLSX]   setup double-checks: Jira key, the client's
          component in shared projects, finance tab, mode, people in Jira vs the tab, role guesses (read only)
  init    --finance XLSX [--tab TAB] [values...] [--write]   draft or refresh the config from the finance tab
  fetch   --month YYYY-MM                                      Tempo + Jira -> records.json
  import  --month YYYY-MM --csv FILE                           a client-format CSV -> records.json (fallback)
  build   --month YYYY-MM --finance XLSX [--adjust F] [--downpayment F] [--overrides F] [--invoiceable F]
          records.json -> changes to apply (cells, new rows/columns, total fixes, notes), report, highlights
  preview --month YYYY-MM --finance XLSX --invoices XLSX       local xlsx copies with the changes (needs openpyxl)
  history --month YYYY-MM [--months 3]                         billable hours per row x person, earlier months
  audit   --month YYYY-MM --finance XLSX [--months 3]          flags total formulas and earlier months (no change)
  plan    --month YYYY-MM --plan FILE [--finance XLSX] [--apply]   Planned cells + Tempo allocations (dry run
          unless --apply)

Auth for Jira/Tempo: env JIRA_PERSONAL_TOKEN (a Jira personal access token; Tempo Server accepts it).
Standard library only (xlsx_read.py reads the finance export). Nothing here writes to Google Sheets: the skill
does that from the JSON this script produces, after the PM's approval.
"""
import sys

sys.dont_write_bytecode = True  # no __pycache__ next to the script

import argparse, calendar, csv, datetime as dt, json, os, re, urllib.error, urllib.parse, urllib.request
from collections import OrderedDict, defaultdict

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import xlsx_read as xl  # stdlib reader next to this script

HERE = os.path.dirname(os.path.abspath(__file__))
MONTHS = list(calendar.month_name)[1:]
DAYS = list(calendar.day_name)
CFG = {}
HUB = None    # the hub root: the nearest folder up from here (or --hub) that holds brain.config.json
SOURCE = None  # where the config lives: ("hub", <brain.config.json>) -> its "finance" section, or ("file", <path>)

# record fields a client CSV column can show (config: client_csv.columns[].field)
CSV_FIELDS = {"epic", "issue_key", "summary", "billed_h", "logged_h", "full_name", "status", "project_name",
              "project_key", "account_key", "account_lead", "account_customer", "started", "comment"}
REQUIRED = ["name", "jira.url", "jira.projects", "jira.home_project", "finance.tab", "people"]


def required(cfg):
    """Workstream mode needs the epic mappings; an invoice needs a rate."""
    req = list(REQUIRED)
    if cfg.get("mode", "workstream") == "workstream":
        req.append("workstreams")
    if get(cfg, "sheets.invoices"):
        req += ["billing.rate", "billing.currency"]
    if get(cfg, "client_csv.filename"):
        req.append("client_csv.columns")
    return req


def project_mode():
    """project mode: one finance row per month for the whole scope (finance.row_label), no epic split."""
    return CFG.get("mode", "workstream") == "project"


def get(d, path, default=None):
    for k in path.split("."):
        if not isinstance(d, dict) or d.get(k) is None:
            return default
        d = d[k]
    return d


def find_hub(start=None):
    d = os.path.abspath(start or os.getcwd())
    while True:
        if os.path.exists(os.path.join(d, "brain.config.json")):
            return d
        parent = os.path.dirname(d)
        if parent == d:
            sys.exit("No brain.config.json in this folder or above: run this from inside the project's hub "
                     "(or pass --hub <path>)")
        d = parent


def set_source(a):
    """The project's config: the hub's brain.config.json > "finance" (one hub = one project), or --config FILE."""
    global HUB, SOURCE
    HUB = find_hub(a.hub)
    SOURCE = ("file", os.path.abspath(a.config)) if a.config else ("hub", os.path.join(HUB, "brain.config.json"))


def read_source():
    kind, path = SOURCE
    if not os.path.exists(path):
        return {}
    data = json.load(open(path, encoding="utf-8"), object_pairs_hook=OrderedDict)
    return data.get("finance") or {} if kind == "hub" else data


def write_source(cfg):
    kind, path = SOURCE
    if kind == "hub":  # keep the rest of brain.config.json as it is
        data = json.load(open(path, encoding="utf-8"), object_pairs_hook=OrderedDict)
        data["finance"] = cfg
        cfg = data
    with open(path, "w", encoding="utf-8") as f:
        json.dump(cfg, f, indent=2, ensure_ascii=False)
        f.write("\n")


def source_name():
    kind, path = SOURCE
    return "brain.config.json > finance" if kind == "hub" else os.path.basename(path)


def load_config():
    cfg = read_source()
    if not cfg:
        sys.exit(f"No finance config in {source_name()}: run `init` first.")
    missing = [k for k in required(cfg) if get(cfg, k) in (None, "", [])]
    if missing:
        sys.exit(f"{source_name()} is missing: {', '.join(missing)} (ask the PM, never guess)")
    bad = [c["field"] for c in get(cfg, "client_csv.columns", []) if c["field"] not in CSV_FIELDS]
    if bad:
        sys.exit(f"client_csv.columns has unknown fields {bad}; allowed: {sorted(CSV_FIELDS)}")
    return cfg


def project_id(cfg=None):
    """Folder name for the work files: config "id", else the project name in lower case."""
    cfg = cfg if cfg is not None else CFG
    return cfg.get("id") or re.sub(r"[^a-z0-9]+", "-", str(cfg.get("name") or "project").lower()).strip("-")


def workdir(month):
    d = os.path.join(HUB, ".work", "finance", project_id(), month)
    os.makedirs(d, exist_ok=True)
    return d


def month_bounds(month):
    y, m = map(int, month.split("-"))
    return dt.date(y, m, 1), dt.date(y, m, calendar.monthrange(y, m)[1])


def shift_month(month, n):
    y, m = map(int, month.split("-"))
    m += n
    y, m = y + (m - 1) // 12, (m - 1) % 12 + 1
    return f"{y}-{m:02d}"


def month_name(month):
    return MONTHS[int(month[5:]) - 1]


def month_label(month):
    """How the finance tab names a month in its month column (default: the English month name)."""
    return get(CFG, "finance.month_label", "{month}").format(month=month_name(month), year=month[:4])


def norm(s):
    return re.sub(r"\s+", " ", str(s or "")).strip()


# ---------------------------------------------------------------- Jira / Tempo (read unless told otherwise)

class Jira:
    def __init__(self, url=None):
        tok = os.environ.get("JIRA_PERSONAL_TOKEN")
        if not tok:
            sys.exit("JIRA_PERSONAL_TOKEN is not set (see .claude/settings.local.example.json)")
        self.base = (url or CFG["jira"]["url"]).rstrip("/")
        self.h = {"Authorization": "Bearer " + tok, "Accept": "application/json",
                  "Content-Type": "application/json"}
        self.cache = {}

    def call(self, path, body=None, method=None):
        req = urllib.request.Request(self.base + path, headers=self.h,
                                     data=None if body is None else json.dumps(body).encode(),
                                     method=method or ("GET" if body is None else "POST"))
        with urllib.request.urlopen(req, timeout=60) as r:
            return json.loads(r.read().decode() or "null")

    def get(self, path):
        if path not in self.cache:
            try:
                self.cache[path] = self.call(path)
            except urllib.error.HTTPError as e:
                self.cache[path] = {"_error": e.code}
        return self.cache[path]


def fetch(month):
    j = Jira()
    start, end = month_bounds(month)
    jc = CFG["jira"]
    # extra scope: shared projects where a component names the client, e.g. OPS + component "Acme"
    extra = {}
    for x in jc.get("extra_scope") or []:
        if x["project"] not in jc["projects"]:
            extra.setdefault(x["project"], []).append(x)
    wls = j.call("/rest/tempo-timesheets/4/worklogs/search",
                 {"from": start.isoformat(), "to": end.isoformat(), "projectKey": jc["projects"] + sorted(extra)})
    fields = j.get("/rest/api/2/field")
    epic_link = next((f["id"] for f in fields if f.get("name") == jc.get("epic_link_field", "Epic Link")), None)
    epic_name = next((f["id"] for f in fields if f.get("name") == jc.get("epic_name_field", "Epic Name")), None)
    flist = ",".join(x for x in ["summary", "status", "project", "parent", "components", epic_link] if x)

    def issue(key):
        return j.get(f"/rest/api/2/issue/{key}?fields={flist}")

    keys = sorted({wl["issue"]["key"] for wl in wls})
    for i in range(0, len(keys), 100):  # one search per 100 issues instead of one call per issue
        res = j.call("/rest/api/2/search", {"jql": f"key in ({','.join(keys[i:i + 100])})",
                                            "fields": flist.split(","), "maxResults": 100})
        for it in res.get("issues", []):
            j.cache[f"/rest/api/2/issue/{it['key']}?fields={flist}"] = it

    def epic_of(key):
        it = issue(key)
        f = it.get("fields", {})
        ek = f.get(epic_link) if epic_link else None
        if not ek and f.get("parent"):
            ek = (issue(f["parent"]["key"]).get("fields", {}) or {}).get(epic_link)
        if not ek:
            return None, None
        ef = j.get(f"/rest/api/2/issue/{ek}?fields=summary" + (f",{epic_name}" if epic_name else "")).get("fields", {})
        return ek, (ef.get(epic_name) if epic_name else None) or ef.get("summary")

    def user(key):
        u = j.get("/rest/api/2/user?key=" + urllib.parse.quote(key))
        return u.get("displayName", key)

    def in_filter(keys):
        """Keys that match the project's Tempo filter (the billing scope); no filter = everything fetched."""
        keys = sorted(set(keys))
        if not jc.get("filter_id"):
            return set(keys)
        ok = set()
        for i in range(0, len(keys), 100):
            jql = f"filter = {jc['filter_id']} AND key in ({','.join(keys[i:i + 100])})"
            res = j.call("/rest/api/2/search", {"jql": jql, "fields": ["key"], "maxResults": 1000})
            ok |= {x["key"] for x in res.get("issues", [])}
        return ok

    def account(wl):
        attr = (wl.get("attributes") or {}).get("_Account_") or {}
        k = wl["issue"].get("accountKey") or attr.get("value")
        if not k:
            return "", "", ""
        a = j.get("/rest/tempo-accounts/1/account/key/" + urllib.parse.quote(k))
        return k, (a.get("lead") or {}).get("displayName", ""), (a.get("customer") or {}).get("name", "")

    def extra_rule(key):
        """The extra_scope entry this issue falls under (by component name), or None."""
        comps = {c.get("name") for c in issue(key).get("fields", {}).get("components") or []}
        return next((x for x in extra.get(key.split("-")[0], []) if comps & set(x.get("components", []))), None)

    allowed = in_filter(wl["issue"]["key"] for wl in wls if wl["issue"]["key"].split("-")[0] not in extra)
    recs = []
    for wl in wls:
        key = wl["issue"]["key"]
        rule = extra_rule(key) if key.split("-")[0] in extra else None
        if key not in allowed and not rule:
            continue
        f = issue(key).get("fields", {})
        ek, en = epic_of(key)
        acc, lead, cust = account(wl)
        recs.append({
            "worklog_id": wl.get("tempoWorklogId") or wl.get("id"),
            "epic_key": ek, "epic": en or "",
            "issue_key": key, "summary": f.get("summary") or wl["issue"].get("summary", ""),
            "billed_h": (wl.get("billableSeconds") or 0) / 3600,
            "logged_h": (wl.get("timeSpentSeconds") or 0) / 3600,
            "full_name": user(wl["worker"]),
            "status": ((f.get("status") or {}).get("name") or "").upper(),
            "project_name": (f.get("project") or {}).get("name", ""),
            "project_key": key.split("-")[0],
            "account_key": acc, "account_lead": lead, "account_customer": cust,
            "started": wl["started"][:16],
            "comment": wl.get("comment") or "",
        })
        if rule:
            recs[-1].update(scope=f'{rule["project"]}/{"+".join(rule["components"])}',
                            scope_include=rule.get("include", "report"))
    return recs


def import_csv(path):
    """Read a CSV in this project's client format (client_csv.columns) back into records."""
    by_field = {c["field"]: c["header"] for c in CFG["client_csv"]["columns"]}
    recs = []
    for r in csv.DictReader(open(path, encoding="utf-8-sig")):
        v = lambda f: r.get(by_field.get(f, ""), "") or ""
        h = float(v("billed_h") or 0)
        rec = {"worklog_id": None, "epic_key": None, "epic": v("epic"), "issue_key": v("issue_key"),
               "summary": v("summary"), "billed_h": h, "logged_h": float(v("logged_h") or h),
               "full_name": v("full_name"), "status": v("status"), "project_name": v("project_name"),
               "project_key": v("issue_key").split("-")[0], "account_key": v("account_key"),
               "account_lead": v("account_lead"), "account_customer": v("account_customer"),
               "started": parse_date(v("started")).strftime("%Y-%m-%d %H:%M"), "comment": v("comment")}
        recs.append(rec)
    return recs


# ---------------------------------------------------------------- build

def num(h):
    """Hours for a note: 8 not 8.0, 7.25 not 7.250000001."""
    return f"{round(float(h), 2):g}"


def fmt_hours(h):
    s = f"{round(h, 9):.9f}".rstrip("0").rstrip(".")
    return s or "0"


DATE_TOKENS = {"YYYY": "{d.year}", "MM": "{d.month:02d}", "M": "{d.month}", "DD": "{d.day:02d}", "D": "{d.day}",
               "HH": "{d.hour:02d}", "H": "{d.hour}", "mm": "{d.minute:02d}"}


def date_template():
    """client_csv.date_format, e.g. "M/D/YYYY H:mm" (no zero padding) -> a str.format template."""
    pat = get(CFG, "client_csv.date_format", "M/D/YYYY H:mm")
    return re.sub("|".join(sorted(DATE_TOKENS, key=len, reverse=True)), lambda m: DATE_TOKENS[m.group(0)], pat)


def fmt_date(s):
    return date_template().format(d=dt.datetime.strptime(s, "%Y-%m-%d %H:%M"))


def parse_date(s):
    pat = get(CFG, "client_csv.date_format", "M/D/YYYY H:mm")
    rx = re.sub("|".join(sorted(DATE_TOKENS, key=len, reverse=True)),
                lambda m: f"(?P<{m.group(0)}>\\d+)", re.escape(pat).replace("\\ ", " "))
    g = re.fullmatch(rx, s.strip()).groupdict()
    pick = lambda *ks: int(next(g[k] for k in ks if k in g))
    return dt.datetime(pick("YYYY"), pick("MM", "M"), pick("DD", "D"), pick("HH", "H"), pick("mm"))


def workstream(epic):
    for w in CFG.get("workstreams", []):
        if epic and re.search(w["match"], epic):
            return w
    return None


def classify(recs, month):
    """Return (kept, flags). Each kept record gets r['ws'] (finance row name), r['ws_cfg'] and r['inv'] (invoice block)."""
    mname = month_name(month)
    home = CFG["jira"]["home_project"]
    other = CFG.get("other_projects") or {}
    default_ws = next((w for w in CFG.get("workstreams", []) if w["finance_row"] == other.get("default_workstream")), None)
    todo = set(get(CFG, "billing.todo_statuses", ["TO DO", "OPEN", "BACKLOG"]))
    account = get(CFG, "jira.tempo_account")
    kept, flags = [], defaultdict(list)
    rules = {f'{x["project"]}/{"+".join(x["components"])}': x.get("include", "report")
             for x in get(CFG, "jira.extra_scope", [])}
    for r in recs:
        if r.get("scope"):  # the config decides, so a rule change needs no new fetch
            r["scope_include"] = rules.get(r["scope"], "report")
        if r.get("scope_include") == "report":  # extra scope shown in the report only, never billed
            flags["extra_report"].append(r)
            continue
        if r["billed_h"] <= 0:
            if r["logged_h"] > 0:
                if project_mode():  # still part of the row's Logged hours
                    r["ws"] = finance_row_label()
                flags["non_billable"].append(r)
            continue
        if r.get("scope_include") == "bill":
            flags["extra_billed"].append(r)
        if project_mode():
            r["ws"], r["ws_cfg"] = finance_row_label(), None
            r["inv"] = get(CFG, "invoice_header", CFG["name"]).format(month=mname, year=month[:4])
            check(r, flags, home, account, todo)
            kept.append(r)
            continue
        w = workstream(r["epic"])
        if r["project_key"] != home and not r["epic"] and not r.get("scope_include"):
            # another project with no epic of ours: the listed people default to one workstream, everyone else is out
            if default_ws and r["full_name"] in other.get("default_people", []):
                r["epic"], w = other.get("default_epic", "{month} " + default_ws["finance_row"]).format(
                    month=mname, year=month[:4]), default_ws
                flags["defaulted"].append(r)
            else:
                flags["other_project_no_epic"].append(r)
                continue
        if w is None:
            flags["unknown_workstream"].append(r)
            r["ws"], r["ws_cfg"] = r["epic"] or "(no epic)", None
        else:
            r["ws"], r["ws_cfg"] = w["finance_row"], w
        # invoice block: one per epic-based header, so two epics on one finance row stay separate lines
        r["inv"] = (w["invoice_header"] if w else "{epic}").format(
            epic=r["epic"] or "(no epic)", month=mname, year=month[:4])
        check(r, flags, home, account, todo)
        kept.append(r)
    return kept, flags


def finance_row_label():
    return get(CFG, "finance.row_label", CFG["name"])


def check(r, flags, home, account, todo):
    if r["project_key"] == home and not r["account_lead"]:
        flags["no_account"].append(r)
    if account and r.get("account_key") and r["account_key"] != account:
        flags["other_account"].append(r)
    if r["status"] in todo:
        flags["todo_status"].append(r)
    if r["billed_h"] > r["logged_h"] + 1e-9:
        flags["billed_gt_logged"].append(r)


def csv_value(r, field):
    if field in ("billed_h", "logged_h"):
        return fmt_hours(r[field])
    if field == "started":
        return fmt_date(r["started"])
    return r.get(field) or ""


def client_csv(kept, path):
    rows = sorted(kept, key=lambda r: r["started"])  # stable: keeps fetch order within a timestamp
    cols = CFG["client_csv"]["columns"]
    with open(path, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, lineterminator="\r\n", quoting=csv.QUOTE_MINIMAL)
        w.writerow([c["header"] for c in cols])
        for r in rows:
            w.writerow([csv_value(r, c["field"]) for c in cols])
    return len(rows), sum(r["billed_h"] for r in rows)


def description(ws, person):
    d = CFG.get("descriptions") or {}
    for table in (d.get(ws, {}), d.get("*", {})):
        if person in table:
            return table[person], table is d.get(ws, {})
        if "*" in table:
            return table["*"], True
    return "TODO: describe", False


def invoice_lines(kept, month, downpayments, overrides=None):
    """overrides: {"<header>": {"invoice": n, "lines": {"<person>": "description"}}} written per run
    (descriptions drafted from each line's task list, invoice split agreed with the PM)."""
    overrides = overrides or {}
    d = CFG.get("descriptions") or {}
    agg, meta, tasks = OrderedDict(), {}, defaultdict(lambda: defaultdict(float))
    for r in kept:
        agg.setdefault(r["inv"], defaultdict(float))[r["full_name"]] += r["billed_h"]
        meta[r["inv"]] = (r["ws"], r["ws_cfg"])
        tasks[(r["inv"], r["full_name"])][f'{r["issue_key"]} {r["summary"]}'] += r["billed_h"]
    groups = defaultdict(list)
    for header, people in agg.items():
        ws, cfg = meta[header]
        ov = overrides.get(header, {})
        lines = []
        for p in sorted(people):
            desc, known = description(header if header in d else ws, p)
            if p in ov.get("lines", {}):
                desc, known = ov["lines"][p], True
            top = sorted(tasks[(header, p)].items(), key=lambda kv: -kv[1])
            lines.append({"person": p, "text": f"{p}: {desc}", "hours": round(people[p], 2), "known_desc": known,
                          "tasks": [f"{round(h, 2)}h {t}" for t, h in top]})
        groups[ov.get("invoice") or (cfg.get("invoice_group", 1) if cfg else 9)].append(
            {"workstream": ws, "rank": CFG["workstreams"].index(cfg) if cfg else 99, "header": header, "lines": lines})
    for dp in downpayments:  # {"workstream","person","hours","label"[,"header","group"]} -> extra lines
        g = next((g for g in groups.values() for b in g if b["workstream"] == dp["workstream"]), None)
        block = {"workstream": dp["workstream"], "header": dp.get("header", dp["workstream"] + " (downpayment)"),
                 "lines": [{"person": dp["person"], "text": f'{dp["person"]}: {dp["label"]}',
                            "hours": dp["hours"], "known_desc": True}]}
        wcfg = next((w for w in CFG.get("workstreams", []) if w["finance_row"] == dp["workstream"]), {})
        groups[dp.get("group") or wcfg.get("invoice_group") or max(groups or [1])].append(block) if g is None \
            else g.append(block)
    order = [sorted(groups[k], key=lambda b: b.get("rank", 99)) for k in sorted(groups)]
    return [{"invoice": i + 1, "blocks": b} for i, b in enumerate(order)]


# ---------------------------------------------------------------- finance tab

def header_labels(ws):
    """Column letter -> header label. With two header rows, a label is "<row 1> / <row 2>" (row 1 carried right)."""
    n = int(get(CFG, "finance.header_rows", 1))
    labels, carry = {}, ""
    for col in range(1, ws.max_column + 1):
        parts = []
        for r in range(1, n + 1):
            v = norm(ws.cell(r, col).value)
            if r == 1 and n > 1:
                carry = v or carry
                v = carry
            if v:
                parts.append(v)
        if parts:
            labels[xl.get_column_letter(col)] = " / ".join(parts)
    labels.update(NEW_COLS)  # columns this run adds (see plan_new_columns)
    return labels


NEW_COLS = {}  # column letter -> header, for people this run adds columns for


def plan_new_columns(kept, xlsx):
    """People who logged in scope but have no columns in the tab get Planned/Logged/Billable columns appended
    after the last used column, named like the other people's headers ("Logged <name>"). Registered in NEW_COLS
    so every cell, total and review line already uses them. Returns the change for the review."""
    ws, _ = finance_sheet(xlsx)
    NEW_COLS.clear()
    labels = header_labels(ws)
    people = people_by_name()
    kinds = [k for k in ("planned", "logged", "billable", "invoiceable")
             if any(k in (p.get("finance_columns") or {}) for p in CFG["people"])]
    need = []
    for name in sorted({r["full_name"] for r in kept}):
        p = people.get(name)
        if p and person_columns(labels, p)[0].get("logged"):
            continue
        need.append(name)
    if not need:
        return []
    last = max(col_index(c) for c in labels)
    groups = [sorted(person_columns(labels, q)[0].values(), key=col_index) for q in CFG["people"]]
    groups = [g for g in groups if g]
    model = max(groups, key=lambda g: col_index(g[-1])) if groups else []  # the right-most person's columns
    words = {"planned": "Planned", "logged": "Logged", "billable": "Billable", "invoiceable": "Invoiceable"}
    out = []
    for name in need:
        p = people.get(name)
        if not p:  # a newcomer: add them to the in-memory config; the PM confirms the role
            p = {"name": name, "role": None, "finance_columns": {}}
            CFG["people"].append(p)
        cols = {}
        for kind in kinds:
            last += 1
            letter = xl.get_column_letter(last)
            header = f"{words[kind]} {p['name']}"
            p["finance_columns"][kind] = header
            NEW_COLS[letter] = header
            cols[kind] = letter
        out.append({"person": name, "role": p.get("role"), "columns": cols,
                    "headers": {k: NEW_COLS[c] for k, c in cols.items()},
                    "copy_format_from": f"{model[0]}:{model[-1]}" if model else None})
    return out


def col_index(letter):
    return xl.column_index_from_string(letter)


def person_columns(labels, person):
    """{"planned"|"logged"|"billable": column letter} for one person, from people[].finance_columns header names.
    A header used by several columns (e.g. a bare "Planned") resolves to the nearest one left of the person's
    Logged column, then Billable."""
    want = person.get("finance_columns") or {}
    found, problems = {}, []
    by_label = defaultdict(list)
    for letter, lab in labels.items():
        by_label[lab].append(letter)
    for kind in ("logged", "billable", "planned", "invoiceable"):
        h = want.get(kind)
        if not h:
            continue
        hits = sorted(by_label.get(norm(h), []), key=col_index)
        if len(hits) > 1:
            anchor = found.get("logged") or found.get("billable")
            left = [c for c in hits if anchor and col_index(c) < col_index(anchor)]
            hits = left[-1:] if left else []
        if hits:
            found[kind] = hits[0]
        else:
            problems.append(f"No '{h}' column for {person['name']} in the finance tab (run `init` again?)")
    return found, problems


def finance_sheet(xlsx):
    wb = xl.load_workbook(xlsx, data_only=True)
    tab = CFG["finance"]["tab"]
    if tab not in wb.sheetnames:
        sys.exit(f"No tab '{tab}' in {xlsx} (tabs: {', '.join(wb.sheetnames)})")
    ws = wb[tab]
    first = int(get(CFG, "finance.header_rows", 1)) + 1
    kc, mc = (col_index(c) for c in get(CFG, "finance.row_key_columns", ["A", "B"]))
    rows = {}
    for row in ws.iter_rows(min_row=first):
        a, b = row[kc - 1].value, row[mc - 1].value
        if a and b:
            rows[(str(a).strip(), str(b).strip())] = row[0].row
    return ws, rows


def people_by_name():
    """Worklog full name -> person; jira_name covers a finance label that differs from the Jira display name."""
    return {p.get("jira_name") or p["name"]: p for p in CFG["people"]}


def finance_updates(kept, month, xlsx, adjust=None, non_billable=()):
    """Map (workstream, person) -> Logged/Billable cells of the finance tab.
    Logged counts the non-billable worklogs too when finance.logged_counts_non_billable (default in project mode,
    as finance tabs filled from a Jira report hold it; set it false where Logged is the time spent of billable worklogs only).
    adjust: [{"person", "workstream"?, "logged"?, "billable"?, "reason"}] overrides a cell with a one-off
    exception agreed with the PM for this month; the reason becomes that cell's note. Without logged/billable
    it only notes the billable cell (e.g. why some logged time is not billed).
    Also returns the Jira-vs-finance comparison of the month rows it touches."""
    ws, rows = finance_sheet(xlsx)
    labels, people = header_labels(ws), people_by_name()
    mlabel = month_label(month)
    agg = defaultdict(lambda: [0.0, 0.0])
    for r in kept:
        agg[(r["ws"], r["full_name"])][0] += r["logged_h"]
        agg[(r["ws"], r["full_name"])][1] += r["billed_h"]
    if get(CFG, "finance.logged_counts_non_billable", project_mode()):
        for r in non_billable:
            if r.get("ws"):
                agg[(r["ws"], r["full_name"])][0] += r["logged_h"]
    cells, problems, new_rows = {}, [], []
    if project_mode() and agg and (finance_row_label(), mlabel) not in rows:
        nr = new_month_row(xlsx, rows, finance_row_label(), mlabel)
        new_rows += [nr] if nr else []
    for (wsn, p), (lg, bl) in sorted(agg.items()):
        row = rows.get((wsn, mlabel))
        cols, _ = person_columns(labels, people[p]) if p in people else ({}, [])
        lc, bc = cols.get("logged"), cols.get("billable")
        if not row:
            problems.append(f"No row '{wsn} / {mlabel}' in the finance tab: add it (with its formulas) first")
            continue
        if not (lc and bc):
            problems.append(f"No Logged/Billable columns for '{p}' in the finance tab")
            continue
        cells[f"{lc}{row}"] = round(lg, 2)
        cells[f"{bc}{row}"] = round(bl, 2)
    notes, jira_cells = {}, dict(cells)
    for adj in adjust or []:
        wsn = adj.get("workstream") or (finance_row_label() if project_mode() else None)
        row = rows.get((wsn, mlabel))
        cols, _ = person_columns(labels, people[adj["person"]]) if adj["person"] in people else ({}, [])
        if not row or not cols:
            problems.append(f"Adjustment for {adj['person']} / {wsn}: no such finance row or columns")
            continue
        if adj.get("logged") is None and adj.get("billable") is None and cols.get("billable"):
            a1, l1 = f"{cols['billable']}{row}", f"{cols.get('logged')}{row}"
            notes[a1] = f"Logged {num(jira_cells.get(l1, 0))}, billable {num(jira_cells.get(a1, 0))}. {adj['reason']}"
            continue
        for kind in ("logged", "billable"):
            if adj.get(kind) is not None and cols.get(kind):
                a1 = f"{cols[kind]}{row}"
                was = jira_cells.get(a1, 0)
                cells[a1] = round(adj[kind], 2)
                notes[a1] = f"Jira {kind}: {num(was)}. {adj['reason']}"
    return (cells, sorted(set(problems)), compare_finance(ws, rows, labels, mlabel, jira_cells, new_rows), notes,
            new_rows)


def compare_finance(ws, rows, labels, mlabel, jira_cells, new_rows=()):
    """Per person and row: what the finance tab holds now vs what Jira gives (the month rows in scope)."""
    names = {finance_row_label()} if project_mode() else {w["finance_row"] for w in CFG.get("workstreams", [])}
    out = []
    for (wsn, m), row in sorted(rows.items(), key=lambda kv: kv[1]):
        if m != mlabel or wsn not in names:
            continue
        for p in CFG["people"]:
            cols, _ = person_columns(labels, p)
            for kind in ("logged", "billable"):
                if not cols.get(kind):
                    continue
                a1 = f"{cols[kind]}{row}"
                added = [n["row"] for n in new_rows]
                old = row - sum(1 for n in added if n <= row)  # the row in the export, before the inserts
                fin, jira = (None if row in added else ws[f"{cols[kind]}{old}"].value), jira_cells.get(a1, 0)
                fin = fin if isinstance(fin, (int, float)) else 0
                if fin or jira:
                    out.append({"row": wsn, "person": p["name"], "kind": kind, "cell": a1,
                                "finance": round(fin, 2), "jira": round(jira, 2), "diff": round(jira - fin, 2)})
    return out


def total_formula(cols, row):
    """SUM rather than a+b+…: a stray text cell (e.g. a "?" typed in a Planned cell) is skipped instead of
    turning the total into #VALUE!."""
    return "=SUM(" + ",".join(f"{c}{row}" for c in cols) + ")"


def new_month_row(xlsx, rows, label, mlabel):
    """Project mode: the month has no row yet. Plan one right under the label's row for the month before (else
    its last month row): A/B labels, every formula of that row shifted to the new row, and C/D/E rebuilt over all
    people's Planned/Logged/Billable columns when the old formula leaves one out. Registers the new row in
    `rows` (rows at or below it move down by one once it is inserted)."""
    wsf = xl.load_workbook(xlsx)[CFG["finance"]["tab"]]
    ws, _ = finance_sheet(xlsx)
    mine = {m: r for (l, m), r in rows.items() if l == label and m in MONTHS}
    if not mine:
        return None
    prev = MONTHS[(MONTHS.index(mlabel) - 1) % 12] if mlabel in MONTHS else None
    anchor = mine.get(prev) or max(mine.values())
    new = anchor + 1
    kc, mc = get(CFG, "finance.row_key_columns", ["A", "B"])
    cells = {f"{kc}{new}": label, f"{mc}{new}": mlabel}
    for c in wsf[anchor]:
        if isinstance(c.value, str) and c.value.startswith("="):
            cells[f"{c.column_letter}{new}"] = re.sub(rf"(?<![A-Z0-9])(\$?[A-Z]{{1,3}}){anchor}(?!\d)",
                                                      rf"\g<1>{new}", c.value)
    cols, fixed = defaultdict(list), []
    for p in CFG["people"]:
        for kind, c in person_columns(header_labels(ws), p)[0].items():
            cols[kind].append(c)
    for total, kind in (("C", "planned"), ("D", "logged"), ("E", "billable")):
        f, want = cells.get(f"{total}{new}", ""), sorted(cols[kind], key=col_index)
        if want and not set(want) <= set(re.findall(r"\$?([A-Z]{1,3})\$?\d+", f)):
            cells[f"{total}{new}"] = total_formula(want, new)
            fixed.append(total)
    for key, r in list(rows.items()):
        if r >= new:
            rows[key] = r + 1
    rows[(label, mlabel)] = new
    return {"insert_after_row": anchor, "row": new, "copy_format_from_row": anchor, "cells": cells,
            "rebuilt_totals": fixed}


def only_new(cells, xlsx, notes=None, new_rows=()):
    """Never change data already in the finance tab: keep only cells that are empty now. A 0 counts as empty
    only in a row with no other figures yet (a fresh month row). Returns (to_write, kept_as_is, equal)."""
    ws, _ = finance_sheet(xlsx)
    people_cols = {c for p in CFG["people"] for c in person_columns(header_labels(ws), p)[0].values()}
    fresh = {}
    write, kept, equal = {}, [], 0
    added = sorted(n["row"] for n in new_rows)
    for a1, v in cells.items():
        col, row = xl.coordinate_from_string(a1)
        if row in added:  # a row this run inserts: everything in it is new
            write[a1] = v
            continue
        row -= sum(1 for n in added if n <= row)  # where that row sits in the export, before the inserts
        a1_now = f"{col}{row}"
        if row not in fresh:
            fresh[row] = not any(isinstance(ws[f"{c}{row}"].value, (int, float)) and ws[f"{c}{row}"].value
                                 for c in people_cols)
        now = ws[a1_now].value
        if now in (None, "") or (now == 0 and fresh[row]):
            write[a1] = v
        elif isinstance(now, (int, float)) and abs(now - v) < 0.005:
            equal += 1
        else:
            kept.append({"cell": a1, "finance": now, "jira": v})
    if notes is not None:
        for a1 in [a for a in notes if a not in write]:
            notes.pop(a1)
    return write, kept, equal


def audit(month, xlsx, months):
    """Flags for the PM to fix separately: month rows whose total formulas (C/D/E) miss a person's column or
    take one of the wrong kind, typed totals, and earlier months whose cells differ from Jira (months whose
    records.json was fetched)."""
    wsf = xl.load_workbook(xlsx)[CFG["finance"]["tab"]]
    ws, rows = finance_sheet(xlsx)
    labels = header_labels(ws)
    cols = defaultdict(set)
    for p in CFG["people"]:
        for kind, c in person_columns(labels, p)[0].items():
            cols[kind].add(c)
    names = {finance_row_label()} if project_mode() else {w["finance_row"] for w in CFG.get("workstreams", [])}
    flags = []
    for (wsn, m), row in sorted(rows.items(), key=lambda kv: kv[1]):
        if wsn not in names or m not in MONTHS:
            continue
        for total, kind in (("C", "planned"), ("D", "logged"), ("E", "billable")):
            f = wsf[f"{total}{row}"].value
            if isinstance(f, str) and f.startswith("="):
                refs = set(re.findall(r"\$?([A-Z]{1,3})\$?\d+", f))
                miss, wrong = sorted(cols[kind] - refs, key=col_index), sorted(refs - cols[kind], key=col_index)
                if miss or wrong:
                    flags.append(f"{wsn} | {m}: {total}{row} `{f}` "
                                 + (f"leaves out {kind} column(s) {', '.join(miss)}" if miss else "")
                                 + ("; " if miss and wrong else "")
                                 + (f"adds column(s) {', '.join(wrong)} that are not {kind}" if wrong else ""))
            elif isinstance(f, (int, float)) and f:
                flags.append(f"{wsn} | {m}: {total}{row} is a typed number ({f}), not a formula")
    data = []
    for k in range(months, 0, -1):
        m = shift_month(month, -k)
        p = os.path.join(workdir(m), "records.json")
        if not os.path.exists(p):
            data.append({"month": m, "skipped": "no records.json (fetch it first)"})
            continue
        kept, fl = classify(json.load(open(p, encoding="utf-8")), m)
        _, _, cmp_, _, _ = finance_updates(kept, m, xlsx, [], fl.get("non_billable", []))
        data += [dict(c, month=m) for c in cmp_ if abs(c["diff"]) >= 0.01]
    return flags, data


def invoiceable(kept, month, xlsx, split=None, reserved=None, new_rows=()):
    """SLA tabs: each month a fixed number of hours is reserved (billing.sla_reserved_hours, e.g. 8 h) and the
    rest is invoiced, so sum(billable) = reserved + sum(invoiceable). The PM decides who carries the reserved
    hours (usually a share of the PM's and DevOps' hours). Without the PM's split this only proposes one:
    the reserved hours shared equally by the roles in preferences.sla_reserved_from (default PM, DevOps),
    each at most their billable hours (e.g. 8 h reserved = 4 h from the PM + 4 h from DevOps)."""
    reserved = float(get(CFG, "billing.sla_reserved_hours") if reserved is None else reserved)
    bill = defaultdict(float)
    for r in kept:
        bill[r["full_name"]] += r["billed_h"]
    people = people_by_name()
    total = round(sum(bill.values()), 2)
    if split is None:
        roles = set(get(CFG, "preferences.sla_reserved_from", ["PM", "DevOps"]))
        pool = [p for p in bill if (people.get(p) or {}).get("role") in roles] or list(bill)
        cut, left = defaultdict(float), reserved
        while left > 1e-6 and any(bill[p] - cut[p] > 1e-6 for p in pool):
            open_ = [p for p in pool if bill[p] - cut[p] > 1e-6]
            for p in open_:
                take = min(left / len(open_), bill[p] - cut[p])
                cut[p] += take
            left = reserved - sum(cut.values())
        split = {p: h - cut[p] for p, h in bill.items()}
    rows, problems = [], []
    for p in sorted(set(bill) | {k for k in split if not k.startswith("_")}):
        inv = round(float(split.get(p, 0)), 2)
        rows.append({"person": p, "role": (people.get(p) or {}).get("role"), "billable": round(bill.get(p, 0), 2),
                     "invoiceable": inv, "reserved": round(bill.get(p, 0) - inv, 2)})
    inv_total = round(sum(r["invoiceable"] for r in rows), 2)
    if abs(total - reserved - inv_total) > 0.02:
        problems.append(f"SLA: billable {total} h - reserved {reserved} h = {round(total - reserved, 2)} h, "
                        f"but the invoiceable split sums to {inv_total} h")
    cells = {}
    if split is not None and from_pm(split):
        ws, rows_by_key = finance_sheet(xlsx)
        labels, row = header_labels(ws), rows_by_key.get((finance_row_label(), month_label(month)))
        row = next((n["row"] for n in new_rows), None) if not row else row + sum(1 for n in new_rows
                                                                                  if n["row"] <= row)
        for r in rows:
            cols, _ = person_columns(labels, people[r["person"]]) if r["person"] in people else ({}, [])
            if not (row and cols.get("invoiceable")):
                problems.append(f"SLA: no Invoiceable column or month row for {r['person']}")
                continue
            cells[f"{cols['invoiceable']}{row}"] = r["invoiceable"]
    return {"reserved": reserved, "billable": total, "invoiceable": inv_total, "rows": rows,
            "cells": cells, "problems": problems, "from_pm": from_pm(split)}


def from_pm(split):
    return bool(split) and split.get("_source") == "pm"


def browse(key):
    return f"{CFG['jira']['url'].rstrip('/')}/browse/{key}"


def row_formula_fixes(xlsx, mlabel, new_rows=()):
    """The month rows this run writes into get correct totals: C/D/E summing every person's Planned/Logged/
    Billable column. A formula that leaves one out, or a typed number, is replaced (listed in the review).
    Inserted rows already get correct totals. Earlier months are only flagged by `audit`."""
    wsf = xl.load_workbook(xlsx)[CFG["finance"]["tab"]]
    ws, rows = finance_sheet(xlsx)
    cols = defaultdict(list)
    for p in CFG["people"]:
        for kind, c in person_columns(header_labels(ws), p)[0].items():
            cols[kind].append(c)
    added = sorted(n["row"] for n in new_rows)
    names = {finance_row_label()} if project_mode() else {w["finance_row"] for w in CFG.get("workstreams", [])}
    fixes = []
    for (label, m), old in rows.items():
        if m != mlabel or label not in names:
            continue
        row = old
        for n in added:  # where the row ends up once the new rows are inserted
            row += 1 if row >= n else 0
        for total, kind in (("C", "planned"), ("D", "logged"), ("E", "billable")):
            want = sorted(cols[kind], key=col_index)
            now = wsf[f"{total}{old}"].value
            refs = set(re.findall(r"\$?([A-Z]{1,3})\$?\d+", now)) if isinstance(now, str) else set()
            if want and not (isinstance(now, str) and now.startswith("=") and set(want) <= refs):
                fixes.append({"cell": f"{total}{row}", "now": now,
                              "new": total_formula(want, row), "kind": kind})
    return fixes


def cell_owner(xlsx):
    """Column letter -> (person, kind) for the review table."""
    ws, _ = finance_sheet(xlsx)
    out = {}
    for p in CFG["people"]:
        for kind, c in person_columns(header_labels(ws), p)[0].items():
            out[c] = (p["name"], kind)
    return out


def highlights(kept, non_billable, month):
    """A few friendly figures for the PM to thank people with. Two are picked per month (rotating)."""
    rate, cur = get(CFG, "billing.rate"), get(CFG, "billing.currency", "")
    per = defaultdict(lambda: {"logged": 0.0, "billed": 0.0, "days": set(), "logs": 0, "described": 0})
    task = defaultdict(lambda: defaultdict(float))
    for r in kept + list(non_billable):
        x = per[r["full_name"]]
        x["logged"] += r["logged_h"]
        x["billed"] += r["billed_h"]
        x["days"].add(r["started"][:10])
        x["logs"] += 1
        x["described"] += 1 if (r.get("comment") or "").strip() else 0
        task[r["issue_key"]][r["full_name"]] += r["billed_h"]
    items = []
    if per:
        p, x = max(per.items(), key=lambda kv: kv[1]["billed"])
        worth = f" ({cur} {x['billed'] * rate:,.0f})" if rate else ""
        items.append(f"🏆 **Most billable hours:** {p}, {round(x['billed'], 1)} h{worth}")
        steady = [(p, x) for p, x in per.items() if x["logged"] >= 10]
        if steady:
            p, x = max(steady, key=lambda kv: (round(kv[1]["billed"] / kv[1]["logged"], 3), kv[1]["logged"]))
            items.append(f"🎯 **Most accurate logs:** {p}, {round(100 * x['billed'] / x['logged'])}% of "
                         f"{round(x['logged'], 1)} logged hours billable")
        p, x = max(per.items(), key=lambda kv: (len(kv[1]["days"]), kv[1]["logged"]))
        items.append(f"📅 **Most steady logger:** {p}, logged on {len(x['days'])} different days")
        wordy = [(p, x) for p, x in per.items() if x["logs"] >= 10]
        if wordy:
            p, x = max(wordy, key=lambda kv: (kv[1]["described"] / kv[1]["logs"], kv[1]["logs"]))
            items.append(f"✍️ **Clearest worklogs:** {p}, {round(100 * x['described'] / x['logs'])}% of "
                         f"{x['logs']} worklogs have a description")
    if task:
        k, who = max(task.items(), key=lambda kv: sum(kv[1].values()))
        summ = next((r["summary"] for r in kept if r["issue_key"] == k), "")[:60]
        items.append(f"🧱 **Biggest task:** {k} {summ} — {round(sum(who.values()), 1)} h "
                     f"({', '.join(p for p, h in sorted(who.items(), key=lambda kv: -kv[1]) if h)})")
    if len(items) <= 2:
        return items, items
    who = lambda s: s.split(":** ", 1)[-1].split(",")[0]  # the person an item is about
    start = int(month.replace("-", "")) % len(items)
    ring = items[start:] + items[:start]
    second = next((i for i in ring[1:] if who(i) != who(ring[0])), ring[1])  # two different people if possible
    return items, [ring[0], second]


def report(month, kept, flags, n, total, inv, cells, problems, compare=(), notes=None, sla=None, held=(),
           equal=0, new_rows=(), fixes=(), owner=None, picked=(), new_cols=()):
    inv_total = sum(l["hours"] for i in inv for b in i["blocks"] for l in b["lines"])
    rate, cur = get(CFG, "billing.rate"), get(CFG, "billing.currency", "")
    money = f" = {cur} {round(inv_total * rate, 2):,.2f}" if rate else " (no rate in the config)"
    invoicing = bool(get(CFG, "sheets.invoices"))
    owner = owner or {}
    L = [f"# Monthly billing — {CFG['name']} — {month}", "",
         f"- {'Client CSV' if get(CFG, 'client_csv.filename') else 'Billable worklogs'}: {n} rows, "
         f"{round(total, 2)} billed hours"]
    L += [f"- Invoice lines: {round(inv_total, 2)} h{money}"] if invoicing else []
    L += [f"- Finance cells to write (empty now): {len(cells)}; already equal: {equal}; "
          f"holding a different value, left as they are: {len(held)}", ""]
    # everything the run will change, in one place, for the PM to read before anything is applied
    L += [f"## Changes to apply in tab '{CFG['finance']['tab']}' (nothing else in the sheet changes)", ""]
    k = 0
    for nc in new_cols:
        k += 1
        L.append(f"{k}. Add columns for {nc['person']} (logged in Jira, no columns yet), after the last column, "
                 f"formatted like {nc['copy_format_from']}: "
                 + ", ".join(f"{c}1 = \"{nc['headers'][kind]}\"" for kind, c in nc["columns"].items())
                 + ("" if nc["role"] else " — role not known yet: ask the PM"))
    for nr in new_rows:
        k += 1
        L.append(f"{k}. Insert row {nr['row']} (under row {nr['insert_after_row']}, same format): "
                 + ", ".join(f"{a1} = `{v}`" for a1, v in nr["cells"].items()))
    for f in fixes:
        k += 1
        L.append(f"{k}. Fix total {f['cell']} ({f['kind']}): `{f['now']}` → `{f['new']}` "
                 "(sums every person's column)")
    if cells:
        k += 1
        L += [f"{k}. Write {len(cells)} empty cells:", "", "   | Cell | Person | Column | Value |",
              "   |---|---|---|---|"]
        for a1, v in cells.items():
            col = re.match(r"[A-Z]+", a1).group(0)
            p, kind = owner.get(col, ("", ""))
            L.append(f"   | {a1} | {p} | {kind} | {v} |")
    for a1, t in (notes or {}).items():
        k += 1
        L.append(f"{k}. Note on {a1}: \"{t}\"")
    if not k:
        L.append("Nothing to change.")
    L.append("")
    per = defaultdict(float)
    for r in kept:
        per[r["full_name"]] += r["billed_h"]
    L += ["## Billable per person", "", "| Person | Hours |", "|---|---|"]
    L += [f"| {p} | {round(h, 2)} |" for p, h in sorted(per.items())] + [""]
    per = defaultdict(float)
    for r in kept:
        per[r["full_name"]] += r["billed_h"]
    L += ["## Billable per person", "", "| Person | Hours |", "|---|---|"]
    L += [f"| {p} | {round(h, 2)} |" for p, h in sorted(per.items())] + [""]
    other = CFG.get("other_projects") or {}
    titles = {
        "unknown_workstream": "Epic not mapped to a workstream (add it to the config)",
        "defaulted": f"Worklogs with no {CFG['name']} epic, put on {other.get('default_workstream')} by default "
                     "(check the title)",
        "other_project_no_epic": f"Other-project worklogs with no {CFG['name']} epic (left out)",
        "todo_status": "Billable time on issues still in TO DO",
        "billed_gt_logged": "Billed hours above logged hours",
        "no_account": f"{CFG['jira']['home_project']} issues without a Tempo account "
                      "(Account Lead/Customer empty in the client CSV)",
        "other_account": f"Worklogs on another Tempo account than {get(CFG, 'jira.tempo_account')}",
        "extra_billed": "Extra-scope worklogs billed (jira.extra_scope include = bill)",
    }
    for k, t in titles.items():
        rs = flags.get(k, [])
        if not rs:
            continue
        L += [f"## {t} ({len(rs)})", ""]
        for r in rs[:40]:
            L.append(f"- {r['started'][:10]} {r['issue_key']} {r['full_name']} "
                     f"billed {round(r['billed_h'], 2)} / logged {round(r['logged_h'], 2)} — {r['summary'][:70]}")
        L.append("")
    nb = defaultdict(lambda: [0.0, 0.0, defaultdict(float)])  # person -> logged, billed, {task: unbilled h}
    for r in kept + flags.get("non_billable", []):
        x = nb[r["full_name"]]
        x[0] += r["logged_h"]
        x[1] += r["billed_h"]
        if r["logged_h"] - r["billed_h"] > 1e-9:
            x[2][f'{r["issue_key"]} {r["summary"][:50]}'.replace("|", "/")] += r["logged_h"] - r["billed_h"]
    nb = {p: x for p, x in nb.items() if x[0] - x[1] >= 0.01}
    if nb:
        L += ["## Logged but not billed (tell the PM; an agreed reason goes into the finance cell's note)", "",
              "| Person | Logged | Billable | Not billed |", "|---|---|---|---|"]
        for p, (lg, bl, t) in sorted(nb.items(), key=lambda kv: kv[1][1] - kv[1][0]):
            L.append(f"| {p} | {round(lg, 2)} | {round(bl, 2)} | {round(lg - bl, 2)} |")
        L.append("")
        for p in sorted(nb, key=lambda q: nb[q][1] - nb[q][0]):
            L += [f"Worklogs of {p} not (fully) billed:", "", "| Date | Issue | Logged | Billable | Task | Worklog |",
                  "|---|---|---|---|---|---|"]
            rs = [r for r in kept + flags.get("non_billable", [])
                  if r["full_name"] == p and r["logged_h"] - r["billed_h"] > 1e-9]
            for r in sorted(rs, key=lambda r: r["started"]):
                L.append(f"| {r['started'][:10]} | [{r['issue_key']}]({browse(r['issue_key'])}) | "
                         f"{round(r['logged_h'], 2)} | {round(r['billed_h'], 2)} | "
                         f"{r['summary'][:60].replace('|', '/')} | {(r.get('comment') or '')[:60].replace('|', '/')} |")
            L.append("")
    if flags.get("extra_report"):
        ex = defaultdict(lambda: [0.0, 0.0])
        for r in flags["extra_report"]:
            ex[(r["scope"], r["full_name"])][0] += r["logged_h"]
            ex[(r["scope"], r["full_name"])][1] += r["billed_h"]
        L += ["## Extra scope, report only (not in the CSV, invoice or finance; the PM decides)", "",
              "| Scope | Person | Logged | Billable |", "|---|---|---|---|"]
        L += [f"| {sc} | {p} | {round(v[0], 2)} | {round(v[1], 2)} |" for (sc, p), v in sorted(ex.items())] + [""]
    if compare:
        diff = [c for c in compare if abs(c["diff"]) >= 0.01]
        L += [f"## Jira vs finance tab ({len(compare) - len(diff)} cells equal, {len(diff)} differ)", ""]
        if diff:
            L += ["| Row | Person | Kind | Cell | Finance now | Jira | Diff |", "|---|---|---|---|---|---|---|"]
            L += [f"| {c['row']} | {c['person']} | {c['kind']} | {c['cell']} | {c['finance']} | {c['jira']} | "
                  f"{c['diff']:+} |" for c in diff]
        L.append("")
    if held:
        L += ["## Not written: the finance tab already holds another value (flag for the PM; fix separately)", "",
              "| Cell | Finance now | Jira |", "|---|---|---|"]
        L += [f"| {h['cell']} | {h['finance']} | {h['jira']} |" for h in held] + [""]
    newdesc = [l["text"] for i in inv for b in i["blocks"] for l in b["lines"] if not l["known_desc"]]
    if newdesc and invoicing:
        L += ["## Invoice descriptions guessed (no history for this workstream + person)", ""]
        L += [f"- {t}" for t in newdesc] + [""]
    if problems:
        L += ["## Finance tab problems", ""] + [f"- {p}" for p in problems] + [""]
    if sla:
        state = "written from the PM's split" if sla["from_pm"] else \
            "PROPOSAL only: the PM distributes it (pass the split with --invoiceable)"
        L += [f"## SLA invoiceable ({state})", "",
              f"Billable {sla['billable']} h = reserved {sla['reserved']} h + invoiceable {sla['invoiceable']} h", "",
              "| Person | Role | Billable | Invoiceable | Reserved share |", "|---|---|---|---|---|"]
        L += [f"| {r['person']} | {r['role'] or ''} | {r['billable']} | {r['invoiceable']} | {r['reserved']} |"
              for r in sla["rows"]] + [""]
    if picked:
        L += ["## Highlights of the month", ""] + [f"> {h}" for h in picked] + [""]
    return "\n".join(L)


# ---------------------------------------------------------------- next-month planning

def history(month, n):
    """Billable hours per workstream x person in the n months before `month`, from their records.json."""
    out, missing = {}, []
    for k in range(n, 0, -1):
        m = shift_month(month, -k)
        p = os.path.join(workdir(m), "records.json")
        if not os.path.exists(p):
            missing.append(m)
            continue
        kept, _ = classify(json.load(open(p, encoding="utf-8")), m)
        agg = defaultdict(float)
        for r in kept:
            agg[f'{r["ws"]} | {r["full_name"]}'] += r["billed_h"]
        out[m] = {k2: round(v, 2) for k2, v in sorted(agg.items())}
    return {"months": out, "missing_records": missing}


def plan_finance_cells(plan, month, xlsx):
    """plan: [{"workstream": epic, "finance_row": row name, "person": name, "hours": h}] -> Planned cells."""
    ws, rows = finance_sheet(xlsx)
    labels, people = header_labels(ws), people_by_name()
    mlabel = month_label(month)
    cells, problems, new_rows = {}, [], []
    if project_mode() and plan and (finance_row_label(), mlabel) not in rows:
        nr = new_month_row(xlsx, rows, finance_row_label(), mlabel)
        new_rows += [nr] if nr else []
    for it in plan:
        it.setdefault("finance_row", finance_row_label() if project_mode() else None)
        row = rows.get((it["finance_row"], mlabel))
        cols, _ = person_columns(labels, people[it["person"]]) if it["person"] in people else ({}, [])
        if not cols.get("planned"):
            problems.append(f"No Planned column for {it['person']} in the finance tab")
            continue
        if not row:
            problems.append(f"Add row '{it['finance_row']} / {mlabel}' (inside its block) before writing the plan")
            continue
        a1 = f"{cols['planned']}{row}"
        cells[a1] = round(cells.get(a1, 0) + it["hours"], 2)
    return cells, sorted(set(problems)), new_rows


def working_days(start, end):
    off = {DAYS.index(d) for d in get(CFG, "billing.weekend_days", ["Saturday", "Sunday"])}
    return sum(1 for d in range((end - start).days + 1) if (start + dt.timedelta(d)).weekday() not in off)


def tempo_plans(plan, month, apply=False):
    """One Tempo allocation per person and workstream on the home project, for the whole month.
    Tempo stores time per working day, so secondsPerDay is set from the total and corrected after a read-back."""
    j = Jira()
    start, end = month_bounds(month)
    proj = j.get(f"/rest/api/2/project/{CFG['jira']['home_project']}")
    people = people_by_name()
    start_time = get(CFG, "tempo_planning.start_time", "09:00")
    out = []
    for it in plan:
        uname = (people.get(it["person"]) or {}).get("jira_username")
        it.setdefault("workstream", "")
        users = j.get("/rest/api/2/user/search?username=" + urllib.parse.quote(uname or it["person"]))
        user = next((u for u in users if (u.get("name") == uname if uname else u.get("displayName") == it["person"])),
                    None) if isinstance(users, list) else None
        total = int(round(it["hours"] * 3600))
        days = working_days(start, end)
        row = {"person": it["person"], "description": it["workstream"], "hours": it["hours"],
               "start": str(start), "end": str(end), "user_key": user and user["key"],
               "working_days": days, "seconds_per_day": max(1, round(total / days))}
        if not user:
            row["error"] = "Jira user not found"
        elif apply:
            # Tempo Server's AllocationBean: assignee by username, the project as plan item, no reviewer
            body = {"assignee": {"key": user["name"], "type": "user"},
                    "planItem": {"id": int(proj["id"]), "type": "PROJECT"},
                    "scope": {"id": int(proj["id"]), "type": "none"},
                    "commitment": 0, "secondsPerDay": row["seconds_per_day"],
                    "includeNonWorkingDays": False, "start": str(start), "startTime": start_time, "end": str(end),
                    "description": it["workstream"], "recurrence": {"rule": "NEVER", "endDate": str(end)}}
            res = j.call("/rest/tempo-planning/1/allocation", body)
            got = j.call(f"/rest/tempo-planning/1/allocation/{res['id']}")
            if abs(got["seconds"] - total) > 60 and got["seconds"]:
                body["secondsPerDay"] = round(body["secondsPerDay"] * total / got["seconds"])
                j.call(f"/rest/tempo-planning/1/allocation/{res['id']}", body, method="PUT")
                got = j.call(f"/rest/tempo-planning/1/allocation/{res['id']}")
            row.update(allocation_id=res["id"], planned_hours=round(got["seconds"] / 3600, 2))
        out.append(row)
    return out


# ---------------------------------------------------------------- invoice tab + local preview

def invoice_tab_name(month):
    nxt = shift_month(month, 1)
    return get(CFG, "sheets.invoice_tab_name", "{next_month} {next_year}").format(
        month=month_name(month), year=month[:4], next_month=month_name(nxt), next_year=nxt[:4])


def invoice_rows(inv):
    """Rows for the invoice tab: [description, hours or None, price or None, formula or None]."""
    rows, rate = [["Description", "Amount", "Price", "Total"]], get(CFG, "billing.rate")
    if rate is None:
        sys.exit("billing.rate is not in the config: ask the PM before building an invoice")
    label = get(CFG, "sheets.invoice_label", "INVOICE {n}")
    for i in inv:
        rows.append([label.format(n=i["invoice"]), None, None, None])
        for b in i["blocks"]:
            rows.append([b["header"], None, None, None])
            for l in b["lines"]:
                n = len(rows) + 1
                rows.append([l["text"], l["hours"], rate, f"=B{n}*C{n}"])
    last = len(rows)
    rows.append(["Total", f"=SUM(B2:B{last})", rate, f"=SUM(D2:D{last})"])
    return rows


def preview(month, wd, finance_xlsx, invoices_xlsx):
    """Apply the outputs to local xlsx copies so they can be checked before anything goes to Google."""
    try:  # optional: only the local preview writes xlsx
        import openpyxl
        from openpyxl.styles import Font
    except ImportError:
        sys.exit("preview needs the optional package openpyxl: pip install openpyxl (nothing else needs it)")
    out = os.path.join(wd, "preview")
    os.makedirs(out, exist_ok=True)
    inv = json.load(open(os.path.join(wd, "invoice.json"), encoding="utf-8"))
    cells = json.load(open(os.path.join(wd, "finance_cells.json"), encoding="utf-8"))
    fwb = openpyxl.load_workbook(finance_xlsx)
    fws = fwb[CFG["finance"]["tab"]]
    for a1, v in cells.items():
        fws[a1].value = v
        fws[a1].font = Font(bold=True, color="C00000")  # preview only: shows what changed
    fwb.save(os.path.join(out, "finance-preview.xlsx"))
    iwb = openpyxl.load_workbook(invoices_xlsx)
    name = invoice_tab_name(month)
    if name in iwb.sheetnames:
        del iwb[name]
    ws = iwb.create_sheet(name, 0)
    money = '"' + get(CFG, "billing.currency_symbol", CFG["billing"]["currency"] + " ") + '"#,##0.00'
    for r in invoice_rows(inv):
        ws.append(r)
        row = ws.max_row
        if r[1] is None and row > 1:
            ws.cell(row, 1).font = Font(bold=True)
        if r[2] is not None:
            ws.cell(row, 3).number_format = ws.cell(row, 4).number_format = money
    ws.column_dimensions["A"].width = 70
    iwb.save(os.path.join(out, "invoices-preview.xlsx"))
    return out


# ---------------------------------------------------------------- init (config draft from the finance tab)

KINDS = {"planned": "planned", "plan": "planned", "logged": "logged", "billable": "billable", "billed": "billable",
         "invoiceable": "invoiceable"}


def discover_people(labels):
    """Read the finance header labels into people with their Planned/Logged/Billable headers.
    Handles "Logged Jane Doe", "BillableJane Doe", "Jane Doe Logged", a bare "Planned" left of "Logged Jane Doe".
    Headers it cannot read (e.g. a bare name, "Invoiceable X") are returned for the PM to place."""
    rx_pre = re.compile(r"^(planned|plan|logged|billable|billed|invoiceable)\s*(.*)$", re.I)
    rx_post = re.compile(r"^(.*?)\s*\b(planned|plan|logged|billable|billed|invoiceable)$", re.I)
    people, order, bare, unread = {}, [], [], []
    for letter in sorted(labels, key=col_index):
        lab = labels[letter].split(" / ")[-1] if " / " in labels[letter] else labels[letter]
        m = rx_pre.match(lab) or rx_post.match(lab)
        if not m:
            unread.append({"column": letter, "header": labels[letter]})
            continue
        kind, name = (m.group(1), m.group(2)) if m.re is rx_pre else (m.group(2), m.group(1))
        if " / " in labels[letter]:
            name = labels[letter].split(" / ")[0]
        name = norm(name)
        if not name:
            bare.append((letter, KINDS[kind.lower()], labels[letter]))
            continue
        if name not in people:
            people[name], _ = {"name": name, "role": None, "finance_columns": {}}, order.append(name)
        people[name]["finance_columns"].setdefault(KINDS[kind.lower()], labels[letter])
        people[name].setdefault("_letters", {})[KINDS[kind.lower()]] = letter
    for letter, kind, lab in bare:  # a bare kind header belongs to the person whose next column follows it
        owner = next((n for n in order if any(col_index(l) == col_index(letter) + 1
                                               for l in people[n]["_letters"].values())), None)
        if owner and kind not in people[owner]["finance_columns"]:
            people[owner]["finance_columns"][kind] = lab
            people[owner]["_letters"][kind] = letter
        else:
            unread.append({"column": letter, "header": lab})
    # a "person" with neither Logged nor Billable is a total column, e.g. "SLA Invoiceable"
    for n in [n for n in order if not {"logged", "billable"} & set(people[n]["finance_columns"])]:
        order.remove(n)
        unread += [{"column": l, "header": h} for h, l in zip(people[n]["finance_columns"].values(),
                                                               people[n]["_letters"].values())]
    return [people[n] for n in order], sorted(unread, key=lambda u: col_index(u["column"]))


def tab_mode(month_rows):
    """workstream when some month has rows under two labels (one block per workstream); one row per month =
    project, even when one month's label differs (e.g. "Acme - Jira log")."""
    return "workstream" if any(sum(1 for k in month_rows if k[1] == m) > 1 for m in MONTHS) else "project"


def init(a):
    base = read_source()
    is_new = not base
    wb = xl.load_workbook(a.finance, data_only=True)
    if not a.tab:
        print(json.dumps({"tabs": wb.sheetnames, "hint": "pass --tab <the project's finance tab>"}, indent=1))
        return
    if a.tab not in wb.sheetnames:
        sys.exit(f"No tab '{a.tab}' (tabs: {', '.join(wb.sheetnames)})")
    global CFG
    CFG = base
    CFG.setdefault("finance", {}).update(tab=a.tab, header_rows=a.header_rows or get(base, "finance.header_rows", 1))
    ws = wb[a.tab]
    labels = header_labels(ws)
    found, unread = discover_people(labels)
    _, rows = finance_sheet(a.finance)
    month_rows = [k for k in rows if k[1] in MONTHS and not k[0].lower().startswith("total")]
    blocks = sorted({k[0] for k in month_rows})
    months =sorted({k[1] for k in rows}, key=lambda m: MONTHS.index(m) if m in MONTHS else 99)

    cfg = OrderedDict(base) if base else OrderedDict([
        ("name", a.name), ("jira", {}), ("billing", {}), ("other_projects", None), ("sheets", {}),
        ("finance", {}), ("client_csv", {}), ("people", []), ("workstreams", []), ("descriptions", {}),
        ("preferences", {})])
    set_ = lambda sec, k, v: v is not None and cfg.setdefault(sec, {}).__setitem__(k, v)
    if a.name:
        cfg["name"] = a.name
    if a.mode or is_new:  # read from the tab
        cfg["mode"] = a.mode or tab_mode(month_rows)
    if cfg.get("mode") == "project" and blocks:  # the label most month rows use
        cfg["finance"].setdefault("row_label", max(blocks, key=lambda b: sum(1 for k in month_rows if k[0] == b)))
    set_("jira", "url", a.jira_url)
    set_("jira", "projects", a.keys and [k.strip() for k in a.keys.split(",")])
    set_("jira", "filter_id", a.filter)
    set_("jira", "home_project", a.home or (a.keys and a.keys.split(",")[0].strip()))
    set_("jira", "tempo_account", a.account)
    set_("billing", "rate", a.rate)
    set_("billing", "currency", a.currency)
    set_("billing", "weekend_days", a.weekend and [d.strip().title() for d in a.weekend.split(",")])
    set_("sheets", "invoices", a.invoices_id)
    set_("sheets", "finance", a.finance_id)
    set_("client_csv", "drive_folder", a.drive_folder)
    cfg["finance"] = {**cfg.get("finance", {}), **CFG["finance"]}
    # people: keep names/roles already confirmed, refresh their headers, add newcomers from the tab
    known = {p["name"]: p for p in cfg.get("people", [])}
    people = []
    for p in found:
        p.pop("_letters", None)
        if p["name"] in known:
            p = {**known.pop(p["name"]), "finance_columns": p["finance_columns"]}
        people.append(p)
    people += [{**p, "_note": "not found in the finance tab any more"} for p in known.values()]
    cfg["people"] = people
    if not cfg.get("client_csv", {}).get("columns"):
        cfg.setdefault("client_csv", {})["columns"] = [
            {"header": h, "field": f} for h, f in [
                ("Epic Link", "epic"), ("Issue Key", "issue_key"), ("Issue summary", "summary"),
                ("Billed Hours", "billed_h"), ("Full name", "full_name"), ("Issue Status", "status"),
                ("Project Name", "project_name"), ("Account Lead", "account_lead"),
                ("Account Customer", "account_customer"), ("Work date", "started"),
                ("Work Description", "comment")]]
    todo = [k for k in required(cfg) if get(cfg, k) in (None, "", [])]
    todo += [f"people[{p['name']}].role" for p in people if not p.get("role")]
    todo += [f"people[{p['name']}].finance_columns.{k}" for p in people for k in ("logged", "billable")
             if k not in p["finance_columns"]]
    cfg["_todo"] = todo
    out = {"config": source_name(), "written": bool(a.write), "todo": todo, "unread_headers": unread,
           "finance_rows": {"workstreams": blocks, "months": months}, "draft": cfg}
    if a.write:
        write_source(cfg)
    if a.check_jira and get(cfg, "jira.url"):
        j = Jira(get(cfg, "jira.url"))
        if get(cfg, "jira.filter_id"):
            f = j.get(f"/rest/api/2/filter/{get(cfg, 'jira.filter_id')}")
            out["jira_filter"] = {"name": f.get("name"), "jql": f.get("jql"), "error": f.get("_error")}
        out["jira_projects"] = {k: j.get(f"/rest/api/2/project/{k}").get("name")
                                for k in get(cfg, "jira.projects", [])}
    print(json.dumps(out, indent=1, ensure_ascii=False))


# ---------------------------------------------------------------- setup double-checks

ROLE_WORDS = {
    "BE": r"\b(be|backend|back-end|api|php|symfony|laravel|migration|endpoint)\b",
    "FE": r"\b(fe|frontend|front-end|ui|angular|react|css|responsive)\b",
    "QA": r"\b(qa|test|testing|regression)\b",
    "PM": r"\b(pm|spec|specification|meeting|client|planning|roadmap|grooming|requirements?|stakeholder|scope)\b",
    "DevOps": r"\b(devops|server|deploy|deployment|live update|docker|backup|nginx|cron|ssl|database export)\b",
    "Design": r"\b(design|figma|ux|wireframe|mockup|prototype)\b",
}


def similar(a, b):
    import difflib
    a, b = re.sub(r"[^a-z0-9]", "", str(a).lower()), re.sub(r"[^a-z0-9]", "", str(b).lower())
    if not a or not b:
        return 0.0
    return 1.0 if a == b else 0.9 if a in b or b in a else round(difflib.SequenceMatcher(None, a, b).ratio(), 2)


def best(name, options, n=3):
    ranked = sorted(((similar(name, o), o) for o in options), key=lambda x: -x[0])
    return [{"value": o, "score": sc} for sc, o in ranked[:n] if sc >= 0.5]


def setup_check(a):
    """What the skill can find on its own, for the PM to confirm with one click: Jira project key, the client's
    component in each shared project (jira.shared_projects / --shared, e.g. a DevOps project with one component
    per client), finance tab, mode, people in Jira vs the finance tab, and a role guess per person."""
    global CFG
    CFG = dict(read_source()) or {"name": a.name}
    name = a.name or CFG.get("name") or sys.exit("check needs --name <the project's name>")
    j = Jira(get(CFG, "jira.url") or a.jira_url or sys.exit("needs jira.url in the config or --jira-url"))
    out = {"name": name}
    projects = {p["key"]: p["name"] for p in j.get("/rest/api/2/project")}
    keys = get(CFG, "jira.projects") or ([a.keys] if a.keys else [])
    out["jira_project"] = ({"configured": {k: projects.get(k, "NOT VISIBLE to this token") for k in keys}} if keys
                           else {"candidates": best(name, projects.values())})
    if not keys and out["jira_project"]["candidates"]:
        top = out["jira_project"]["candidates"][0]["value"]
        keys = [k for k, v in projects.items() if v == top]
        out["jira_project"]["key_of_best"] = keys[0]
    shared = get(CFG, "jira.shared_projects") or (a.shared.split(",") if a.shared else [])
    comp_ids, out["shared_projects"] = set(), {}
    for sp in shared:
        comps = {c["id"]: c["name"] for c in j.get(f"/rest/api/2/project/{sp}/components") or []
                 if isinstance(c, dict)}
        have = [x for x in get(CFG, "jira.extra_scope", []) if x["project"] == sp]
        pick = have[0]["components"] if have else \
            [c["value"] for c in (best(name, comps.values()) or best(projects.get(keys[0], "") if keys else "",
                                                                        comps.values()))[:1]]
        out["shared_projects"][sp] = {"configured": pick} if have else \
            {"candidates": best(name, comps.values()), "all": sorted(comps.values())}
        comp_ids |= {i for i, n in comps.items() if n in pick}
    fin_people, tab = [], None
    if a.finance:
        wb = xl.load_workbook(a.finance, data_only=True)
        tab = get(CFG, "finance.tab")
        out["finance_tab"] = {"configured": tab, "exists": tab in wb.sheetnames} if tab else \
            {"candidates": best(name, wb.sheetnames)}
        tab = tab if tab in wb.sheetnames else (out["finance_tab"].get("candidates") or [{}])[0].get("value")
        if tab:
            CFG.setdefault("finance", {}).setdefault("tab", tab)
            ws, rows = finance_sheet(a.finance)
            mrows = [k for k in rows if k[1] in MONTHS and not k[0].lower().startswith("total")]
            out["mode_guess"] = {"mode": tab_mode(mrows), "row_labels": sorted({k[0] for k in mrows})}
            fin_people = [p["name"] for p in discover_people(header_labels(ws))[0]]
    # people and roles from the last months' worklogs (scope keys + the client's components in shared projects)
    end = dt.date.today().replace(day=1) - dt.timedelta(days=1)
    start = (end.replace(day=1) - dt.timedelta(days=31 * (a.months - 1))).replace(day=1)
    wls = j.call("/rest/tempo-timesheets/4/worklogs/search",
                 {"from": str(start), "to": str(end), "projectKey": keys + shared}) if keys else []
    name_of = lambda k: j.get("/rest/api/2/user?key=" + urllib.parse.quote(k)).get("displayName", k)
    hours, workers = defaultdict(lambda: [0.0, 0.0]), {}
    for w in wls:
        i = w["issue"]
        if i["projectKey"] in shared and not comp_ids & {str(c) for c in i.get("components") or []}:
            continue
        u = workers.setdefault(w["worker"], name_of(w["worker"]))
        hours[u][0] += (w.get("timeSpentSeconds") or 0) / 3600
        hours[u][1] += (w.get("billableSeconds") or 0) / 3600
    # role guess: keywords over everything these people logged (all projects), a project's role from
    # preferences.role_by_project (e.g. {"OPS": "DevOps"}), plus Tempo team memberships
    by_project = get(CFG, "preferences.role_by_project", {})
    roles = defaultdict(lambda: defaultdict(float))
    allw = j.call("/rest/tempo-timesheets/4/worklogs/search",
                  {"from": str(start), "to": str(end), "worker": sorted(workers)}) if workers else []
    for w in allw:
        u, h, i = workers.get(w["worker"]), (w.get("timeSpentSeconds") or 0) / 3600, w["issue"]
        if not u:
            continue
        if i["projectKey"] in by_project:
            roles[u][by_project[i["projectKey"]]] += h
        for role, rx in ROLE_WORDS.items():
            if re.search(rx, i.get("summary", "").lower()):
                roles[u][role] += h
    teams = defaultdict(list)
    for t in j.get("/rest/tempo-teams/2/team") if isinstance(j.get("/rest/tempo-teams/2/team"), list) else []:
        mem = j.get(f"/rest/tempo-teams/2/team/{t['id']}/member")
        if isinstance(mem, list) and len(mem) < 30:  # skip company-wide teams
            for m in mem:
                role = ((m.get("membership") or {}).get("role") or {}).get("name")
                teams[norm(m["member"].get("displayname"))].append(
                    t["name"] + (f": {role}" if role and role != "Member" else ""))
    out["period"] = f"{start} .. {end}"
    out["people"] = []
    for u, (lg, bl) in sorted(hours.items(), key=lambda kv: -kv[1][0]):
        rs = sorted(roles[u].items(), key=lambda kv: -kv[1])
        tot = sum(roles[u].values()) or 1
        out["people"].append({"name": u, "logged": round(lg, 2), "billable": round(bl, 2),
                              "role_guess": [f"{r} {round(100 * h / tot)}%" for r, h in rs[:2]],
                              "tempo_teams": teams.get(u, []),
                              "finance_column": u in fin_people if tab else None})
    if tab:
        out["finance_people_without_logs"] = [p for p in fin_people if p not in hours]
    print(json.dumps(out, indent=1, ensure_ascii=False))


def main():
    global CFG
    ap = argparse.ArgumentParser()
    ap.add_argument("cmd", choices=["init", "check", "fetch", "import", "build", "preview", "history", "plan",
                                    "audit"])
    ap.add_argument("--hub", help="the hub folder (default: the nearest folder up that holds brain.config.json)")
    ap.add_argument("--config", help="a config JSON file instead of the hub's brain.config.json > finance")
    ap.add_argument("--month", help="YYYY-MM of the work being invoiced (plan: the month being planned)")
    ap.add_argument("--csv")
    ap.add_argument("--finance", help="xlsx export of the finance sheet")
    ap.add_argument("--invoices", help="xlsx export of the invoices sheet (preview)")
    ap.add_argument("--downpayment", help="JSON list of extra downpayment lines")
    ap.add_argument("--overrides", help="JSON {header: {invoice: n, lines: {person: description}}}")
    ap.add_argument("--invoiceable", help='SLA: the PM split {"_source": "pm", "<person>": hours}')
    ap.add_argument("--reserved", type=float, help="SLA: reserved hours this month (default billing.sla_reserved_hours)")
    ap.add_argument("--adjust", help="JSON [{person, workstream?, logged?, billable?, reason}] finance overrides")
    ap.add_argument("--plan", help="plan JSON: [{workstream, finance_row, person, hours}] (plan)")
    ap.add_argument("--apply", action="store_true", help="plan: create the Tempo allocations (default: dry run)")
    ap.add_argument("--months", type=int, default=3, help="history: how many earlier months")
    g = ap.add_argument_group("init")
    g.add_argument("--tab"), g.add_argument("--header-rows", type=int), g.add_argument("--name")
    g.add_argument("--mode", choices=["project", "workstream"], help="project: one row per month; workstream: by epic")
    g.add_argument("--jira-url"), g.add_argument("--keys", help="comma-separated Jira project keys")
    g.add_argument("--filter", type=int, help="Tempo/Jira filter id = billing scope"), g.add_argument("--home")
    g.add_argument("--account", help="Tempo account key"), g.add_argument("--rate", type=float)
    g.add_argument("--currency"), g.add_argument("--weekend", help="e.g. Friday,Saturday")
    g.add_argument("--invoices-id"), g.add_argument("--finance-id"), g.add_argument("--drive-folder")
    g.add_argument("--write", action="store_true", help="init: write the config (default: print the draft)")
    g.add_argument("--shared", help="check: comma-separated shared project keys (one component per client)")
    g.add_argument("--check-jira", action="store_true", help="init: read the filter and project names from Jira")
    a = ap.parse_args()
    set_source(a)
    if a.cmd == "check":
        return setup_check(a)
    if a.cmd == "init":
        if not a.finance:
            sys.exit("init needs --finance <xlsx export of the finance workbook>")
        return init(a)
    if not a.month:
        sys.exit("--month YYYY-MM is required")
    CFG = load_config()
    wd = workdir(a.month)
    rec_path = os.path.join(wd, "records.json")
    if a.cmd in ("fetch", "import"):
        recs = fetch(a.month) if a.cmd == "fetch" else import_csv(a.csv)
        json.dump(recs, open(rec_path, "w", encoding="utf-8"), indent=1, ensure_ascii=False)
        print(f"{len(recs)} worklogs -> {rec_path}")
        return
    if a.cmd == "audit":
        if not a.finance:
            sys.exit("audit needs --finance <xlsx>")
        fl, data = audit(a.month, a.finance, a.months)
        L = [f"# Finance tab audit — {CFG['name']} — before {a.month}", "",
             "Flags only: nothing here is changed by the skill; the PM fixes them separately if they want.", ""]
        L += [f"## Total formulas ({len(fl)})", ""] + [f"- {f}" for f in fl] + [""]
        diff = [d for d in data if "diff" in d]
        L += [f"## Earlier months: finance vs Jira ({len(diff)} cells differ)", ""]
        L += [f"- {d['month']} — skipped: {d['skipped']}" for d in data if "skipped" in d]
        if diff:
            L += ["", "| Month | Row | Person | Kind | Cell | Finance | Jira | Diff |", "|---|---|---|---|---|---|---|---|"]
            L += [f"| {d['month']} | {d['row']} | {d['person']} | {d['kind']} | {d['cell']} | {d['finance']} | "
                  f"{d['jira']} | {d['diff']:+} |" for d in diff]
        txt = "\n".join(L) + "\n"
        open(os.path.join(wd, "audit.md"), "w", encoding="utf-8").write(txt)
        print(txt)
        return
    if a.cmd == "history":
        print(json.dumps(history(a.month, a.months), indent=1, ensure_ascii=False))
        return
    if a.cmd == "plan":
        plan = json.load(open(a.plan, encoding="utf-8"))
        res = {"month": a.month, "total_hours": round(sum(i["hours"] for i in plan), 2)}
        if a.finance:
            cells, res["problems"], res["new_rows"] = plan_finance_cells(plan, a.month, a.finance)
            res["finance_cells"], res["not_written"], _ = only_new(cells, a.finance, None, res["new_rows"])
        res["tempo"] = tempo_plans(plan, a.month, a.apply)
        res["applied"] = a.apply
        json.dump(res, open(os.path.join(wd, "plan_result.json"), "w", encoding="utf-8"), indent=1, ensure_ascii=False)
        print(json.dumps(res, indent=1, ensure_ascii=False))
        return
    if a.cmd == "preview":
        print("preview ->", preview(a.month, wd, a.finance, a.invoices))
        return
    recs = json.load(open(rec_path, encoding="utf-8"))
    kept, flags = classify(recs, a.month)
    start, end = month_bounds(a.month)
    if get(CFG, "client_csv.filename"):  # no client CSV for projects that don't send one
        n, total = client_csv(kept, os.path.join(wd, CFG["client_csv"]["filename"].format(**{"from": start, "to": end})))
    else:
        n, total = len(kept), sum(r["billed_h"] for r in kept)
    dps = json.load(open(a.downpayment, encoding="utf-8")) if a.downpayment else []
    ovr = json.load(open(a.overrides, encoding="utf-8")) if a.overrides else {}
    inv = invoice_lines(kept, a.month, dps, ovr)
    adj = json.load(open(a.adjust, encoding="utf-8")) if a.adjust else []
    new_cols = plan_new_columns(kept, a.finance) if a.finance else []
    json.dump(new_cols, open(os.path.join(wd, "finance_new_columns.json"), "w", encoding="utf-8"), indent=1,
              ensure_ascii=False)
    cells, problems, cmp_, notes, new_rows = (
        finance_updates(kept, a.month, a.finance, adj, flags.get("non_billable", []))
        if a.finance else ({}, [], [], {}, []))
    sla = None
    if a.finance and get(CFG, "billing.sla_reserved_hours") is not None:
        split = json.load(open(a.invoiceable, encoding="utf-8")) if a.invoiceable else None
        sla = invoiceable(kept, a.month, a.finance, split, a.reserved, new_rows)
        cells.update(sla["cells"])
        problems += sla["problems"]
    json.dump(inv, open(os.path.join(wd, "invoice.json"), "w", encoding="utf-8"), indent=1, ensure_ascii=False)
    # finance_cells_all.json = what Jira gives; finance_cells.json = only the cells that are empty in the tab now
    json.dump(cells, open(os.path.join(wd, "finance_cells_all.json"), "w", encoding="utf-8"), indent=1)
    held, equal = [], 0
    if a.finance:
        cells, held, equal = only_new(cells, a.finance, notes, new_rows)
    json.dump(cells, open(os.path.join(wd, "finance_cells.json"), "w", encoding="utf-8"), indent=1)
    # rows to insert before writing the cells (row numbers in finance_cells.json already count them)
    json.dump(new_rows, open(os.path.join(wd, "finance_new_rows.json"), "w", encoding="utf-8"), indent=1,
              ensure_ascii=False)
    fixes = row_formula_fixes(a.finance, month_label(a.month), new_rows) if a.finance else []
    json.dump(fixes, open(os.path.join(wd, "finance_formula_fixes.json"), "w", encoding="utf-8"), indent=1,
              ensure_ascii=False)
    if a.finance:
        json.dump({"compare": cmp_, "notes": notes, "not_written": held}, open(os.path.join(
            wd, "finance_compare.json"), "w", encoding="utf-8"), indent=1, ensure_ascii=False)
    every, picked = highlights(kept, flags.get("non_billable", []), a.month)
    json.dump({"all": every, "picked": picked}, open(os.path.join(wd, "highlights.json"), "w", encoding="utf-8"),
              indent=1, ensure_ascii=False)
    rep = report(a.month, kept, flags, n, total, inv, cells, problems, cmp_, notes, sla, held, equal, new_rows,
                 fixes, cell_owner(a.finance) if a.finance else {}, picked, new_cols)
    open(os.path.join(wd, "report.md"), "w", encoding="utf-8").write(rep)
    print(rep)
    print(f"\nFiles in {wd}")


if __name__ == "__main__":
    main()
